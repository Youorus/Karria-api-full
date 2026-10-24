"""Spécification + vue → DOCX : styles nommés, puces Word, ornements ancrés, polices incorporées."""

from __future__ import annotations

import copy
import io
import re
import zipfile
from typing import Any

import pytest
from docx import Document
from fontTools.ttLib import TTFont

from documents.infrastructure.fakes import (
    CENTERED_ROLES,
    CV_ROLES,
    SERVED,
    centered_template_pdf,
    cv_font_set,
    cv_template_pdf,
    labels_for,
    png,
)
from karria.modules.documents.domain.view import from_cv, from_sample
from karria.modules.documents.infrastructure.docx_renderer import (
    DocxRenderer,
    FontNames,
    anchored_shape,
    embed_fonts,
    emu,
    obfuscate,
    rgb,
    set_widths,
    twips,
)
from karria.modules.documents.infrastructure.fonts import FontFile, FontSet
from karria.modules.documents.infrastructure.pdf_reader import PdfDocument
from karria.modules.documents.infrastructure.spec_builder import SpecBuilder

pytestmark = pytest.mark.unit

Json = dict[str, Any]
KEY = "{0123ABCD-4567-89EF-0123-456789ABCDEF}"


def _build(pdf: bytes, roles: dict[str, tuple[str, str, int]], shape: str | None) -> tuple:
    with PdfDocument(pdf) as doc:
        page = doc.read()
        labels = labels_for(page.lines, roles)
        labels["photo_shape"] = shape
        return SpecBuilder(doc, page, labels, cv_font_set(), SERVED).build()


@pytest.fixture(scope="module")
def cv_built() -> tuple:
    return _build(cv_template_pdf(), CV_ROLES, "circle")


@pytest.fixture(scope="module")
def centered_built() -> tuple:
    return _build(centered_template_pdf(), CENTERED_ROLES, None)


@pytest.fixture(scope="module")
def fonts() -> FontSet:
    return cv_font_set()


@pytest.fixture
def built(cv_built: tuple) -> tuple:
    return copy.deepcopy(cv_built)


CV: Json = {
    "identity": {"first_name": "Ada", "last_name": "Lovelace", "headline": "Ingénieure"},
    "contact": {
        "phone": "06 00 00 00 00",
        "email": "ada@example.org",
        "city": "Londres",
        "links": [{"url": "https://ada.dev"}],
    },
    "summary": "Mathématicienne passionnée par les machines et la poésie des nombres.",
    "experiences": [
        {
            "position": "Analyste",
            "company": "Babbage",
            "start_date": "2020",
            "end_date": "2023",
            "missions": ["Machine analytique", "Notes célèbres"],
        },
        {"position": "Stagiaire", "company": "Royal", "start_date": "2019"},
    ],
    "educations": [{"degree": "Mathématiques", "school": "Maison", "start_date": "2015"}],
    "skills": [{"name": "Python", "category": "Tech"}, {"name": "Écoute", "category": "Humain"}],
    "languages": [{"language": "Français", "level": "Natif"}],
    "interests": [{"name": "Poésie"}],
}


def unzip(docx: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(docx)) as z:
        return {n: z.read(n) for n in z.namelist()}


def render(
    spec: Json, view: Json, fonts: FontSet, assets: dict[str, bytes] | None = None, **kw: Any
) -> bytes:
    return DocxRenderer(spec, view, fonts, **{"photo": None, **kw}).render(assets or {})


# ---------------------------------------------------------------------- petites fonctions


class TestOutils:
    def test_les_unites_word(self) -> None:
        assert str(rgb("#a1b2c3")) == "A1B2C3"
        assert emu(1) == 36000
        assert twips(25.4) == 1440

    def test_lobfuscation_appliquee_deux_fois_rend_la_police(self) -> None:
        data = bytes(range(64))
        once = obfuscate(data, KEY)
        assert once != data
        assert once[32:] == data[32:]
        assert obfuscate(once, KEY) == data

    def test_lobfuscation_suit_la_cle_inversee_ecma(self) -> None:
        once = obfuscate(bytes(32), KEY)
        hexkey = KEY.strip("{}").replace("-", "")
        assert once[:16] == bytes.fromhex(hexkey)[::-1]

    def test_les_noms_word_suivent_la_famille_ancienne(self, fonts: FontSet) -> None:
        names = FontNames(fonts)
        assert names.word_font("Testa", 700, False) == ("Testa", True, False)
        assert names.word_font("Testa", 400, False) == ("Testa", False, False)
        assert names.word_font("Absente", 600, True) == ("Absente", True, True)

    def test_un_filet_devient_un_rectangle_plein(self) -> None:
        drawing = anchored_shape(
            x_mm=10,
            y_mm=5,
            w_mm=0,
            h_mm=20,
            relative_h="page",
            relative_v="page",
            kind="vline",
            fill=None,
            stroke="#123456",
            stroke_pt=0.1,
            behind=False,
        )
        xml = re.sub(r">\s+<", "><", drawing.xml)
        assert 'behindDoc="0"' in xml
        assert 'prst="rect"' in xml
        assert '<a:srgbClr val="123456"/>' in xml
        assert "<a:ln><a:noFill/></a:ln>" in xml

    def test_une_courbe_tracee_est_une_ellipse_contour(self) -> None:
        xml = anchored_shape(
            x_mm=0,
            y_mm=0,
            w_mm=-1,
            h_mm=2,
            relative_h="column",
            relative_v="line",
            kind="curve",
            fill=None,
            stroke="#ABCDEF",
            stroke_pt=1.0,
        ).xml
        assert 'prst="ellipse"' in xml
        assert "<a:noFill/>" in xml
        assert 'w="12700"' in xml
        assert 'cx="0"' in xml

    def test_un_filet_horizontal_est_centre_sur_son_trait(self) -> None:
        xml = anchored_shape(
            x_mm=0,
            y_mm=10,
            w_mm=30,
            h_mm=0,
            relative_h="page",
            relative_v="paragraph",
            kind="hline",
            fill="#000000",
            stroke=None,
            stroke_pt=0.75,
        ).xml
        assert f"<wp:posOffset>{emu(10 - 0.75 * 25.4 / 72 / 2)}</wp:posOffset>" in xml


# ---------------------------------------------------------------------- document complet


class TestDocument:
    @pytest.fixture
    def docx(self, built: tuple, fonts: FontSet) -> dict[str, bytes]:
        spec, _, assets = built
        return unzip(render(spec, from_cv(spec, CV), fonts, assets, photo=png(60, 80)))

    def test_un_style_word_nomme_par_role(self, docx: dict[str, bytes]) -> None:
        styles = docx["word/styles.xml"].decode()
        for name in (
            "Titre de section",
            "Intitulé (experiences)",
            "Puce",
            "Compétence",
            "Poste visé",
        ):
            assert f'w:val="Karria · {name}"' in styles

    def test_les_puces_sont_de_vraies_listes_word(self, docx: dict[str, bytes]) -> None:
        numbering = docx["word/numbering.xml"].decode()
        assert 'w:abstractNumId="900"' in numbering
        assert 'w:abstractNumId="901"' in numbering
        assert 'w:lvlText w:val="•"' in numbering
        assert 'w:lvlText w:val="■"' in numbering
        assert '<w:numId w:val="900"/>' in docx["word/document.xml"].decode()

    def test_les_polices_sont_incorporees_et_obfusquees(self, docx: dict[str, bytes]) -> None:
        table = docx["word/fontTable.xml"].decode()
        assert "w:embedRegular" in table
        assert "w:embedBold" in table
        rels = docx["word/_rels/fontTable.xml.rels"].decode()
        assert 'Target="fonts/font1.odttf"' in rels
        assert 'Extension="odttf"' in docx["[Content_Types].xml"].decode()
        assert "embedTrueTypeFonts" in docx["word/settings.xml"].decode()
        key = re.search(r'r:id="rIdKarriaFont1" w:fontKey="(\{[^"]+\})"', table)
        assert key is not None
        font = TTFont(io.BytesIO(obfuscate(docx["word/fonts/font1.odttf"], key.group(1))))
        assert font["name"].getDebugName(1) == "Testa"

    def test_ornements_ancres_et_photo_flottante(self, docx: dict[str, bytes]) -> None:
        xml = docx["word/document.xml"].decode()
        assert xml.count("<wp:anchor") >= 10
        assert 'relativeFrom="line"' in xml
        assert 'relativeFrom="paragraph"' in xml
        assert "<pic:pic" in xml
        assert 'relativeHeight="2"' in xml
        assert "wp:inline" not in xml

    def test_len_tete_a_deux_piles_passe_par_un_tableau(self, docx: dict[str, bytes]) -> None:
        xml = docx["word/document.xml"].decode()
        assert xml.index("Ada Lovelace") < xml.index("06 00 00 00 00")
        assert xml.count("<w:tbl>") == 3

    def test_les_proprietes_portent_le_nom(self, docx: dict[str, bytes]) -> None:
        core = docx["docProps/core.xml"].decode()
        assert "CV — Ada Lovelace" in core
        assert "Ingénieure" in core

    def test_lexemple_garde_la_position_de_sa_photo(self, built: tuple, fonts: FontSet) -> None:
        spec, sample, assets = built
        spec["header"]["photo"]["shape"] = "rect"
        spec["header"]["decorations"].append(
            {"kind": "rect", "x_mm": -2, "y_mm": -3, "w_mm": 20, "h_mm": 10, "fill": "#FF0000"}
        )
        out = unzip(
            render(spec, from_sample(spec, sample), fonts, assets, photo=png(), k=0.9, f=0.96)
        )
        xml = out["word/document.xml"].decode()
        assert '<a:srgbClr val="FF0000"/>' in xml
        assert f'<wp:extent cx="{emu(18)}" cy="{emu(7)}"/>' in xml

    def test_un_modele_a_une_colonne_centre_le_poste_vise(
        self, centered_built: tuple, fonts: FontSet
    ) -> None:
        spec, sample, assets = copy.deepcopy(centered_built)
        view = from_sample(spec, sample)
        view["header"]["headline"] = "Chimiste"
        out = unzip(render(spec, view, fonts, assets, photo=png()))
        xml = out["word/document.xml"].decode()
        assert '<w:jc w:val="center"/>' in xml
        assert '<w:jc w:val="both"/>' in xml
        assert "CENTRES INTERETS" in xml
        assert xml.count("<w:tbl>") == 1

    def test_sans_en_tete_les_formes_sancrent_sur_le_premier_espaceur(
        self, built: tuple, fonts: FontSet
    ) -> None:
        spec, _, assets = built
        view = from_cv(spec, {"skills": [{"name": "Python"}]})
        renderer = DocxRenderer(spec, view, fonts, photo=None)
        renderer.render(assets)
        assert renderer.anchor_paragraph is None
        assert renderer.first_spacer is not None
        assert "<wp:anchor" in renderer.first_spacer._p.xml

    def test_un_document_vide_reste_valide(self, built: tuple, fonts: FontSet) -> None:
        spec, _, assets = built
        data = render(spec, {"header": {}, "sections": []}, fonts, assets)
        assert Document(io.BytesIO(data)).paragraphs == []


# ---------------------------------------------------------------------- en-tête


class TestEnTete:
    def test_les_piles_ecartees_mais_disjointes_restent_un_seul_flux(self) -> None:
        a = {"x_mm": 10, "y_mm": 10}
        b = {"x_mm": 100, "y_mm": 30}
        assert DocxRenderer._header_clusters([a, b]) == [[a, b]]
        assert DocxRenderer._header_clusters([a]) == [[a]]
        c, d = {"x_mm": 100, "y_mm": 5}, {"x_mm": 110, "y_mm": 15}
        assert DocxRenderer._header_clusters([a, {**a, "y_mm": 20}, c, d]) == [
            [a, {**a, "y_mm": 20}],
            [c, d],
        ]

    def test_deux_champs_sur_une_ligne_passent_par_un_taquet(
        self, built: tuple, fonts: FontSet
    ) -> None:
        spec, _, assets = built
        elements = spec["header"]["elements"]
        phone = next(e for e in elements if e["field"] == "phone")
        phone["x_mm"] = 40.0
        phone["y_mm"] = elements[0]["y_mm"] + 1
        phone.pop("icon")
        out = unzip(render(spec, from_cv(spec, CV), fonts, assets))
        xml = out["word/document.xml"].decode()
        assert "<w:tabs><w:tab " in xml
        assert "<w:tab/>" in xml

    def test_des_lignes_qui_se_chevauchent_resserrent_linterligne(
        self, built: tuple, fonts: FontSet
    ) -> None:
        spec, _, assets = built
        elements = spec["header"]["elements"]
        headline = next(e for e in elements if e["field"] == "headline")
        headline["y_mm"] = elements[0]["y_mm"] + 3
        view = from_cv(spec, CV)
        view["header"]["phone"] = ""
        renderer = DocxRenderer(spec, view, fonts, photo=None)
        renderer.render(assets)
        name = renderer.styles["header.headline"]
        paras = [p for p in renderer.doc.paragraphs if p.text == "Ingénieure"]
        assert paras
        assert paras[0].paragraph_format.space_before == 0
        assert paras[0].paragraph_format.line_spacing < renderer.pitch(name) * 36000

    def test_un_pictogramme_sans_fichier_est_ignore(self, built: tuple, fonts: FontSet) -> None:
        spec, _, _ = built
        out = unzip(render(spec, from_cv(spec, CV), fonts, {}))
        assert not [n for n in out if n.startswith("word/media/")]


# ---------------------------------------------------------------------- rubriques


class TestRubriques:
    @pytest.fixture
    def renderer(self, built: tuple, fonts: FontSet) -> DocxRenderer:
        spec, _, assets = built
        r = DocxRenderer(spec, from_cv(spec, CV), fonts, photo=None)
        r.spec = {**r.spec, "_assets": assets}
        return r

    def test_un_style_absent_prend_celui_des_titres(self, renderer: DocxRenderer) -> None:
        assert renderer.style("x.inconnu") == "Karria · inconnu"
        assert renderer.lead("x.inconnu") == renderer.lead("section_title")
        assert renderer.style("autre.inconnu") == "Karria · inconnu"

    def test_un_interlettrage_passe_dans_le_style_et_le_run(self, renderer: DocxRenderer) -> None:
        renderer.styles["section_title"] = {
            **renderer.styles["section_title"],
            "letter_spacing_pt": 1.5,
        }
        renderer.style("section_title")
        assert (
            'w:spacing w:val="30"' in renderer.doc.styles["Karria · Titre de section"].element.xml
        )
        p = renderer.doc.add_paragraph()
        run = p.add_run("x")
        renderer.apply_run_style(run, "section_title")
        assert 'w:val="30"' in run._r.xml
        renderer.apply_run_style(p.add_run("y"), "absent")

    def test_lespace_avant_dun_paragraphe_sans_style_connu(self, renderer: DocxRenderer) -> None:
        p = renderer.doc.add_paragraph()
        assert renderer._para_lead(p) == renderer.lead("section_title")
        assert renderer._space_before_mm(p) == 0.0
        p.paragraph_format.space_before = 36000
        assert renderer._space_before_mm(p) == pytest.approx(1.0, abs=0.01)

    def test_un_ornement_sur_un_paragraphe_vide_y_est_ajoute(self, renderer: DocxRenderer) -> None:
        p = renderer.doc.add_paragraph()
        renderer.ornament(p, {"kind": "rect", "w_mm": 2}, 0, 0, 2)
        assert "<wp:anchor" in p._p.xml
        renderer.item_ornaments({"decorations": []}, [], last_item=True, col_w=50)

    def test_le_retour_a_la_ligne_compte_les_mots(self, renderer: DocxRenderer) -> None:
        assert renderer.wrap("un deux trois quatre cinq six", "summary.summary", 15) > 1
        assert renderer.wrap("court", "header.full_name", 200) == 1

    def test_une_puce_sans_police_servie_garde_son_retrait(self, renderer: DocxRenderer) -> None:
        renderer.styles["skills.skill"] = {**renderer.styles["skills.skill"], "family": "Absente"}
        p = renderer.doc.add_paragraph()
        mk = {"shape": "circle", "dx_mm": -3.0, "color": "#000000"}
        renderer.bullet(p, "skills.skill", mk, 4.0)
        renderer.bullet(
            renderer.doc.add_paragraph(style=renderer.style("skills.skill")),
            "skills.skill",
            mk,
            4.0,
        )
        assert len(renderer._numbering) == 1
        assert "w:numPr" in p._p.xml

    def test_une_rubrique_complete_avec_variantes(self, built: tuple, fonts: FontSet) -> None:
        spec, sample, assets = built
        col = spec["bands"][0]["columns"][1]
        summary, exp = col["sections"]
        summary["title_w_mm"] = 20.0
        exp["decorations"].append(
            {
                "kind": "vline",
                "w_mm": 0,
                "h_mm": 5,
                "stroke": "#010101",
                "anchor": "section",
                "dx_mm": -1,
                "dy_mm": 0,
                "stretch": True,
                "bottom_inset_mm": 0.0,
            }
        )
        exp["decorations"].append(
            {
                "kind": "rect",
                "w_mm": 1,
                "h_mm": 1,
                "fill": "#020202",
                "anchor": "section",
                "dx_mm": 0,
                "dy_mm": 0,
            }
        )
        exp["decorations"].append(
            {
                "kind": "curve",
                "w_mm": 1,
                "h_mm": 1,
                "fill": "#030303",
                "anchor": "item",
                "dx_mm": 0,
                "dy_mm": 0,
                "attach": "row",
                "ref_role": "absent",
                "center_offset_mm": 0,
            }
        )
        exp["decorations"].append(
            {
                "kind": "rect",
                "w_mm": 1,
                "h_mm": 1,
                "fill": "#040404",
                "anchor": "item",
                "dx_mm": 0,
                "dy_mm": 0,
            }
        )
        exp["slots"][0]["separator"] = None
        exp["slots"][0]["gap_mm"] = None
        view = from_sample(spec, sample)
        view["sections"].append(
            {"kind": "experiences", "title": "Bis", "items": [{"item_dates": "2001"}, {}]}
        )
        xml = unzip(render(spec, view, fonts, assets))["word/document.xml"].decode()
        for color in ("010101", "020202", "030303", "040404"):
            assert f'val="{color}"' in xml
        assert "<w:ind w:right=" in xml or 'w:right="' in xml

    def test_le_dernier_element_sans_trait_propre_na_pas_de_frise(
        self, built: tuple, fonts: FontSet
    ) -> None:
        spec, sample, assets = built
        frise = spec["bands"][0]["columns"][1]["sections"][1]["decorations"][0]
        avec = unzip(render(spec, from_sample(spec, sample), fonts, assets))["word/document.xml"]
        frise["last_bottom_inset_mm"] = None
        sans = unzip(render(spec, from_sample(spec, sample), fonts, assets))["word/document.xml"]
        assert avec.count(b"Karria ornement") == sans.count(b"Karria ornement") + 1

    def test_une_colonne_sans_rubrique_garde_un_paragraphe(
        self, built: tuple, fonts: FontSet
    ) -> None:
        spec, _, assets = built
        view = from_cv(spec, {"summary": "Bonjour"})
        doc = Document(io.BytesIO(render(spec, view, fonts, assets)))
        cells = doc.tables[0].rows[0].cells
        assert cells[1].paragraphs
        assert cells[3].paragraphs[0].text == "Profil"

    def test_ladresse_web_sajoute_aux_coordonnees(self, built: tuple, fonts: FontSet) -> None:
        spec, _, assets = built
        xml = unzip(render(spec, from_cv(spec, CV), fonts, assets))["word/document.xml"].decode()
        assert "ada.dev" in xml
        assert "Londres" in xml


def test_une_police_sans_sous_famille_est_reguliere() -> None:
    from documents.infrastructure.fakes import build_font  # noqa: PLC0415

    data = build_font(family="Solo")
    font = FontFile("Solo", 500, False, data)
    names = FontNames(FontSet([font]))
    assert names.word_font("Solo", 500, False) == ("Solo", False, False)


def test_un_tableau_peut_etre_mis_en_retrait() -> None:
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    set_widths(table, [10.0, 20.0], indent_mm=5.0)
    assert f'w:w="{twips(5.0)}"' in table._tbl.xml
    assert "w:tblInd" in table._tbl.xml


def test_incorporer_deux_fois_ne_duplique_ni_type_ni_reglage(fonts: FontSet) -> None:
    buffer = io.BytesIO()
    Document().save(buffer)
    names = FontNames(fonts)
    twice = unzip(embed_fonts(embed_fonts(buffer.getvalue(), names), names))
    assert twice["[Content_Types].xml"].count(b'Extension="odttf"') == 1
    assert twice["word/settings.xml"].count(b"embedTrueTypeFonts") == 1


def test_une_numerotation_vide_recoit_sa_premiere_definition(built: tuple, fonts: FontSet) -> None:
    from docx.oxml.ns import qn  # noqa: PLC0415

    spec, _, _ = built
    renderer = DocxRenderer(spec, {"header": {}, "sections": []}, fonts, photo=None)
    numbering = renderer.doc.part.numbering_part.element
    for num in numbering.findall(qn("w:num")):
        numbering.remove(num)
    st = renderer.styles["skills.skill"]
    assert renderer.numbering("•", st, "#000000", size_pt=9, left_mm=4, hanging_mm=2) == 900
    assert numbering.find(qn("w:abstractNum")) is not None


def test_un_blanc_deja_a_lecart_nest_pas_elargi() -> None:
    p = Document().add_paragraph()
    DocxRenderer._spaced_run(p, " ", 0.01)
    assert "w:spacing" not in p._p.xml
