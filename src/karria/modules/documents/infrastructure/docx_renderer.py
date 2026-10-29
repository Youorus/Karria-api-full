"""Spécification + vue → DOCX modifiable (python-docx), depuis la **même** spécification que le PDF.

Contrat
-------
- **Un style Word nommé par rôle** (« Karria · Intitulé », « Karria · Puce »…) :
  la personne change une police ou une couleur partout d'un geste.
- **Espacements fidèles** : marges de page, interligne exact, espace avant
  chaque paragraphe, retraits, taquets, largeurs de colonnes, interlettrage
  (`w:spacing` en vingtièmes de point) — tous tirés de la spécification, aux
  mêmes facteurs d'ajustement (`k` espacements, `f` corps) que le PDF.
- Word et LibreOffice posent l'excédent d'un interligne « exact » **au-dessus**
  du texte, Chromium le centre : l'espace avant est compensé d'un demi-blanc.
- Colonnes = cellules d'un tableau sans bordure à largeurs fixes (grille
  **et** cellules : LibreOffice lit l'une, Word l'autre).
- **Ornements** en formes DrawingML ancrées (filets, pastilles, cadres,
  frises) : l'en-tête et la page ancrés à la page, les frises et pastilles
  d'élément ancrées à leur paragraphe — elles suivent le texte quand il bouge.
  La photo est une image flottante ancrée à la page, au cadre du modèle.
- **Polices incorporées** (instances statiques, réduites au latin, obfusquées
  selon ECMA-376) : le DOCX s'affiche avec la vraie police chez une personne
  qui ne l'a pas installée, et LibreOffice (aperçu, mesure) aussi.
"""

from __future__ import annotations

import io
import itertools
import re
import uuid
import zipfile
from typing import Any, Final
from xml.sax.saxutils import quoteattr

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from fontTools.ttLib import TTFont
from lxml import etree

from karria.modules.documents.infrastructure.fonts import (
    FontFile,
    FontSet,
    glyph_box,
    subset_for_documents,
)
from karria.modules.documents.infrastructure.html_renderer import (
    LIST_KINDS,
    ROLE_MARKER,
    estimate,
    extra_contact_slots,
    placement,
    style_key,
)
from karria.shared.images import crop_to_frame

Json = dict[str, Any]

ROLE_FR: Final[dict[str, str]] = {
    "section_title": "Titre de section",
    "item_title": "Intitulé",
    "item_subtitle": "Sous-titre",
    "item_dates": "Dates",
    "item_location": "Lieu",
    "bullet": "Puce",
    "item_text": "Texte",
    "summary": "Profil",
    "skill": "Compétence",
    "language": "Langue",
    "interest": "Intérêt",
    "subsection_title": "Sous-titre de liste",
    "phone": "Téléphone",
    "email": "E-mail",
    "address": "Adresse",
    "website": "Site",
    "first_name": "Prénom",
    "last_name": "Nom",
    "full_name": "Nom complet",
    "headline": "Poste visé",
}
EMU_PER_MM: Final = 36000
MM_TO_PT: Final = 72 / 25.4
W_NS: Final = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS: Final = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
FONT_REL: Final = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"
FONT_CT: Final = "application/vnd.openxmlformats-officedocument.obfuscatedFont"
_SLOTS: Final = {
    (False, False): "embedRegular",
    (True, False): "embedBold",
    (False, True): "embedItalic",
    (True, True): "embedBoldItalic",
}


def rgb(hex_color: str) -> RGBColor:
    """`#RRGGBB` → couleur Word."""
    return RGBColor.from_string(hex_color.lstrip("#")[:6].upper())


def emu(mm: float) -> int:
    """Millimètres → EMU (unités DrawingML)."""
    return round(mm * EMU_PER_MM)


def twips(mm: float) -> int:
    """Millimètres → vingtièmes de point."""
    return round(mm / 25.4 * 1440)


def no_borders(table: Any) -> None:
    """Tableau de mise en page : sans bordures, sans marges de cellule, largeurs fixes."""
    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    pr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), "0")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    pr.append(mar)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    pr.append(layout)


def set_widths(table: Any, widths_mm: list[float], indent_mm: float = 0.0) -> None:
    """Largeurs dans la grille (`w:gridCol`) **et** dans chaque cellule ; retrait du tableau."""
    table.autofit = False
    for i, w in enumerate(widths_mm):
        table.columns[i].width = Mm(w)
        table.rows[0].cells[i].width = Mm(w)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(twips(sum(widths_mm))))
    tw.set(qn("w:type"), "dxa")
    table._tbl.tblPr.append(tw)
    if indent_mm:
        ind = OxmlElement("w:tblInd")
        ind.set(qn("w:w"), str(twips(indent_mm)))
        ind.set(qn("w:type"), "dxa")
        table._tbl.tblPr.append(ind)


class FontNames:
    """Les noms sous lesquels Word connaît chaque police servie.

    Word range une police par **famille ancienne** (`name` ID 1) et quatre
    emplacements (normal, gras, italique, gras italique) : un Raleway 500
    statique s'appelle « Raleway Medium » (normal), un Raleway 700 « Raleway »
    (gras). Le run porte donc ce nom et ce drapeau, pas la graisse CSS.
    """

    def __init__(self, fonts: FontSet) -> None:
        self.by_style: dict[tuple[str, int, bool], tuple[str, bool, bool]] = {}
        self.files: dict[tuple[str, bool, bool], FontFile] = {}
        for f in fonts.files:
            font = TTFont(io.BytesIO(f.data), lazy=True)
            family = str(font["name"].getDebugName(1) or f.family)
            sub = str(font["name"].getDebugName(2) or "Regular").lower()
            bold, italic = "bold" in sub, "italic" in sub
            self.by_style[(f.family, f.weight, f.italic)] = (family, bold, italic)
            self.files.setdefault((family, bold, italic), f)
        self._fonts = fonts

    def word_font(self, family: str, weight: int, italic: bool) -> tuple[str, bool, bool]:
        """(nom Word, gras, italique) pour un style de la spécification."""
        f = self._fonts.get(family, weight, italic)
        if f is not None and (f.family, f.weight, f.italic) in self.by_style:
            return self.by_style[(f.family, f.weight, f.italic)]
        return family, weight >= 600, italic  # noqa: PLR2004


def obfuscate(data: bytes, key: str) -> bytes:
    """Obfuscation ECMA-376 d'une police incorporée : 32 premiers octets ⊕ la clé GUID inversée."""
    hexkey = key.strip("{}").replace("-", "")
    key_bytes = bytes(int(hexkey[i : i + 2], 16) for i in range(30, -1, -2))
    head = bytes(b ^ key_bytes[i % 16] for i, b in enumerate(data[:32]))
    return head + data[32:]


def embed_fonts(docx: bytes, names: FontNames) -> bytes:
    """Incorpore les polices servies dans le DOCX (table des polices, parties obfusquées)."""
    src = zipfile.ZipFile(io.BytesIO(docx))
    files = {n: src.read(n) for n in src.namelist()}
    ns = {"w": W_NS, "r": R_NS}
    table = etree.fromstring(files["word/fontTable.xml"])
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    rels = etree.Element(f"{{{rel_ns}}}Relationships", nsmap={None: rel_ns})
    content = etree.fromstring(files["[Content_Types].xml"])
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    if not any(el.get("Extension") == "odttf" for el in content):
        default = etree.SubElement(content, f"{{{ct_ns}}}Default")
        default.set("Extension", "odttf")
        default.set("ContentType", FONT_CT)
    families: dict[str, Any] = {}
    for font in table.findall("w:font", ns):
        families[font.get(f"{{{W_NS}}}name")] = font
    for n, ((family, bold, italic), file) in enumerate(sorted(names.files.items()), 1):
        node = families.get(family)
        if node is None:
            node = etree.SubElement(table, f"{{{W_NS}}}font")
            node.set(f"{{{W_NS}}}name", family)
            families[family] = node
        key = "{" + str(uuid.uuid4()).upper() + "}"
        rid = f"rIdKarriaFont{n}"
        slot = etree.SubElement(node, f"{{{W_NS}}}{_SLOTS[(bold, italic)]}")
        slot.set(f"{{{R_NS}}}id", rid)
        slot.set(f"{{{W_NS}}}fontKey", key)
        rel = etree.SubElement(rels, f"{{{rel_ns}}}Relationship")
        rel.set("Id", rid)
        rel.set("Type", FONT_REL)
        rel.set("Target", f"fonts/font{n}.odttf")
        files[f"word/fonts/font{n}.odttf"] = obfuscate(subset_for_documents(file.data), key)
    files["word/fontTable.xml"] = etree.tostring(
        table, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    files["word/_rels/fontTable.xml.rels"] = etree.tostring(
        rels, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    files["[Content_Types].xml"] = etree.tostring(
        content, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    settings = etree.fromstring(files["word/settings.xml"])
    if settings.find("w:embedTrueTypeFonts", ns) is None:
        settings.insert(0, etree.Element(f"{{{W_NS}}}embedTrueTypeFonts"))
    files["word/settings.xml"] = etree.tostring(
        settings, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in files.items():
            dst.writestr(name, data)
    return out.getvalue()


_shape_ids = itertools.count(1)


def anchored_shape(
    *,
    x_mm: float,
    y_mm: float,
    w_mm: float,
    h_mm: float,
    relative_h: str,
    relative_v: str,
    kind: str,
    fill: str | None,
    stroke: str | None,
    stroke_pt: float,
    behind: bool = True,
) -> Any:
    """Une forme DrawingML flottante (`wp:anchor` + `wps:wsp`) : filet, pastille, cadre.

    `relative_h`/`relative_v` : `page`, `column`, `paragraph`, `margin`.
    """
    sid = next(_shape_ids)
    w, h = max(w_mm, 0.0), max(h_mm, 0.0)
    line_color = stroke or fill
    if kind in ("vline", "hline"):
        # Un filet devient un rectangle plein de l'épaisseur du trait : une
        # forme « ligne » d'étendue nulle disparaît dans LibreOffice.
        thickness = max(stroke_pt, 0.25) * 25.4 / 72
        if kind == "vline":
            x_mm, w = x_mm - thickness / 2, thickness
        else:
            y_mm, h = y_mm - thickness / 2, thickness
        kind, fill, stroke = "rect", line_color, None
    geometry = "ellipse" if kind == "curve" else "rect"
    fill_xml = (
        f'<a:solidFill><a:srgbClr val="{fill.lstrip("#")}"/></a:solidFill>'
        if fill
        else "<a:noFill/>"
    )
    line_xml = (
        f'<a:ln w="{round(stroke_pt * 12700)}"><a:solidFill><a:srgbClr val="{stroke.lstrip("#")}"/></a:solidFill></a:ln>'
        if stroke
        else "<a:ln><a:noFill/></a:ln>"
    )
    xml = (
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        f'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{sid}" behindDoc="{1 if behind else 0}" '
        'locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="{relative_h}"><wp:posOffset>{emu(x_mm)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="{relative_v}"><wp:posOffset>{emu(y_mm)}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{emu(w)}" cy="{emu(h)}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        f'<wp:docPr id="{sid}" name="Karria ornement {sid}"/><wp:cNvGraphicFramePr/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        "<wps:wsp><wps:cNvSpPr/>"
        f'<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{emu(w)}" cy="{emu(h)}"/></a:xfrm>'
        f'<a:prstGeom prst="{geometry}"><a:avLst/></a:prstGeom>{fill_xml}{line_xml}</wps:spPr>'
        "<wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor>"
    )
    drawing = OxmlElement("w:drawing")
    drawing.append(parse_xml(xml))
    return drawing


class DocxRenderer:
    """Rend une vue dans un modèle, en DOCX modifiable."""

    def __init__(
        self,
        spec: Json,
        view: Json,
        fonts: FontSet,
        *,
        photo: bytes | None,
        k: float = 1.0,
        f: float = 1.0,
    ) -> None:
        """
        Args:
            photo: Octets de la photo (PNG/JPEG), ou `None`.
            k, f: Facteurs d'espacement et de corps retenus par la mise en page du PDF.
        """
        self.spec, self.view, self.fonts = spec, view, fonts
        self.photo, self.k, self.f = photo, k, f
        self.styles: dict[str, Json] = spec["styles"]
        self.names = FontNames(fonts)
        self.doc = Document()
        self._word_styles: dict[str, str] = {}
        self._numbering: dict[tuple[Any, ...], int] = {}
        #: Par paragraphe : espace avant (mm), style, texte, retrait — de quoi
        #: calculer sa hauteur comme Word la posera (lignes × interligne exact).
        self._geometry: dict[int, tuple[float, str, str, float]] = {}
        self.header_bottom = 0.0
        self.anchor_paragraph: Paragraph | None = None
        self.first_spacer: Paragraph | None = None

    # ------------------------------------------------------------------ styles nommés
    def style(self, key: str) -> str:
        """Le style Word nommé d'un style de la spécification (créé à la demande)."""
        if key in self._word_styles:
            return self._word_styles[key]
        st = self.styles.get(key) or self.styles["section_title"]
        role = key.rsplit(".", maxsplit=1)[-1]
        prefix = key.split(".", maxsplit=1)[0] if "." in key else ""
        qualified = (
            prefix and prefix != "header" and role in ("item_title", "item_dates", "item_subtitle")
        )
        name = f"Karria · {ROLE_FR.get(role, role)}" + (f" ({prefix})" if qualified else "")
        if name not in [s.name for s in self.doc.styles]:
            s = self.doc.styles.add_style(name, 1)
            font = s.font
            family, bold, italic = self.names.word_font(
                st["family"], int(st["weight"]), bool(st["italic"])
            )
            font.name = family
            rfonts = s.element.rPr.rFonts
            for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rfonts.set(qn(attr), family)
            scale = 1.0 if key.startswith("header.") else self.f
            font.size = Pt(round(st["size_pt"] * scale * 2) / 2)
            font.bold = bold
            font.italic = italic
            font.color.rgb = rgb(st["color"])
            font.all_caps = bool(st["uppercase"])
            if st["letter_spacing_pt"]:
                sp = OxmlElement("w:spacing")
                sp.set(qn("w:val"), str(round(st["letter_spacing_pt"] * 20)))
                s.element.rPr.append(sp)
            pf = s.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Mm(self.pitch(st) * scale)
            pf.widow_control = False
        self._word_styles[key] = name
        return name

    @staticmethod
    def pitch(st: Json) -> float:
        """Interligne d'un style (mm)."""
        return float(st["pitch_mm"] or st["box_h_mm"])

    def lead(self, key: str) -> float:
        """Blanc d'interligne : Word/LibreOffice le posent **au-dessus** du texte."""
        st = self.styles.get(key) or self.styles["section_title"]
        return self.pitch(st) - float(st["box_h_mm"])

    def para(
        self,
        container: Any,
        key: str,
        text: str,
        *,
        before_mm: float,
        indent_mm: float = 0.0,
        align: str = "left",
    ) -> Paragraph:
        """Un paragraphe au style du rôle, avec l'espace avant mesuré."""
        p: Paragraph = container.add_paragraph(style=self.style(key))
        k = 1.0 if key.startswith("header.") else self.k
        before = before_mm * k - self.lead(key)
        if before < 0:
            # Lignes qui se chevauchent dans le modèle (un nom en très grand
            # corps sous un prénom) : l'espace avant ne peut pas être négatif.
            # On réduit l'interligne exact de ce paragraphe — Word et LibreOffice
            # retirent l'excédent au-dessus du texte, là où il n'y a que le
            # blanc des accents — dans la limite d'un quart de la boîte.
            st = self.styles.get(key) or self.styles["section_title"]
            scale = 1.0 if key.startswith("header.") else self.f
            cut = min(-before, 0.25 * float(st["box_h_mm"]))
            p.paragraph_format.line_spacing = Mm(self.pitch(st) * scale - cut)
            before = 0.0
        p.paragraph_format.space_before = Mm(before)
        p.paragraph_format.left_indent = Mm(indent_mm)
        self._geometry[id(p._p)] = (before, key, text, indent_mm)
        if align == "justify":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(text)
        return p

    # ------------------------------------------------------------------ en-tête et page
    def page_shapes(self, anchor: Paragraph) -> None:
        """Ornements fixes (en-tête, page) et photo : ancrés à la page, derrière le texte."""
        h = self.spec["header"]
        decorations = list(h["decorations"]) + [
            d for d in self.spec["page_decorations"] if d.get("anchor") == "page"
        ]
        run = anchor.add_run()
        for d in decorations:
            x, y, w, hh = d["x_mm"], d["y_mm"], d["w_mm"], d["h_mm"]
            if d["kind"] == "rect" and x < 0:
                w, x = w + x, 0.0
            if d["kind"] == "rect" and y < 0:
                hh, y = hh + y, 0.0
            run._r.append(
                anchored_shape(
                    x_mm=x,
                    y_mm=y,
                    w_mm=w,
                    h_mm=hh,
                    relative_h="page",
                    relative_v="page",
                    kind=d["kind"],
                    fill=d.get("fill"),
                    stroke=d.get("stroke"),
                    stroke_pt=float(d.get("stroke_pt") or 0.75),
                )
            )
        p = h.get("photo")
        if p and self.photo:
            focus = (0.5, 0.3)
            position = p.get("sample_object_position", "")
            if self.view.get("sample") and re.fullmatch(r"-?[\d.]+% -?[\d.]+%", position):
                fx, fy = (min(1.0, max(0.0, float(v.rstrip("%")) / 100)) for v in position.split())
                focus = (fx, fy)
            data = crop_to_frame(
                self.photo,
                ratio=p["w_mm"] / max(p["h_mm"], 0.1),
                focus=focus,
                circle=p.get("shape") == "circle",
            )
            picture = run.add_picture(io.BytesIO(data), width=Mm(p["w_mm"]), height=Mm(p["h_mm"]))
            self._float_picture(run, picture, p["x_mm"], p["y_mm"])

    @staticmethod
    def _float_picture(run: Any, picture: Any, x_mm: float, y_mm: float) -> None:
        """Change une image en ligne en image flottante, à sa place sur la page."""
        inline = run._r.xpath(".//wp:inline")[-1]
        extent = inline.find(qn("wp:extent"))
        graphic = inline.find(qn("a:graphic"))
        doc_pr = inline.find(qn("wp:docPr"))
        anchor = parse_xml(
            '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="2" behindDoc="0" locked="0" '
            'layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="page"><wp:posOffset>{emu(x_mm)}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="page"><wp:posOffset>{emu(y_mm)}</wp:posOffset></wp:positionV>'
            "</wp:anchor>"
        )
        anchor.append(extent)
        anchor.append(
            parse_xml(
                '<wp:effectExtent xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" l="0" t="0" r="0" b="0"/>'
            )
        )
        anchor.append(
            parse_xml(
                '<wp:wrapNone xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"/>'
            )
        )
        anchor.append(doc_pr)
        anchor.append(
            parse_xml(
                '<wp:cNvGraphicFramePr xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"/>'
            )
        )
        anchor.append(graphic)
        inline.getparent().replace(inline, anchor)
        del picture

    def header(self) -> None:
        """L'en-tête : chaque texte à sa hauteur mesurée, dans l'ordre de lecture.

        Les éléments d'une même hauteur à des abscisses différentes passent par
        des taquets de tabulation. Quand l'en-tête a deux piles indépendantes
        (le nom à gauche, les coordonnées à droite, à des hauteurs qui se
        chevauchent), chaque pile va dans une cellule d'un tableau sans
        bordure : le texte reste dans le flux, lisible par un ATS, et chaque
        pile garde ses propres écarts.
        """
        h = self.spec["header"]
        texts = [e for e in h["elements"] if self.view["header"].get(e["field"])]
        page_top = min([e["y_mm"] for e in texts], default=10.0)
        clusters = self._header_clusters(texts)
        if len(clusters) == 1:
            bottom = self._header_flow(self.doc, clusters[0], page_top, 0.0)
        else:
            lefts = [min(e["x_mm"] for e in c) for c in clusters]
            edges = [
                0.0,
                *(lefts[i] - 2.0 for i in range(1, len(lefts))),
                self.spec["page"]["width_mm"],
            ]
            widths = [edges[i + 1] - edges[i] for i in range(len(clusters))]
            table = self.doc.add_table(rows=1, cols=len(clusters))
            no_borders(table)
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            set_widths(table, widths)
            bottom = page_top
            for i, cluster in enumerate(clusters):
                cell = table.rows[0].cells[i]
                cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
                bottom = max(bottom, self._header_flow(cell, cluster, page_top, edges[i]))
        self.header_bottom = bottom
        self.page_top = page_top

    @staticmethod
    def _header_clusters(texts: list[Json]) -> list[list[Json]]:
        """Les piles de l'en-tête, de gauche à droite (abscisses écartées de plus de 25 mm).

        Une seule pile si les piles ne se chevauchent pas en hauteur : le flux
        simple (avec taquets) les rend alors fidèlement.
        """
        groups: list[list[Json]] = []
        for e in sorted(texts, key=lambda e: e["x_mm"]):
            if groups and e["x_mm"] - min(x["x_mm"] for x in groups[-1]) <= 25:  # noqa: PLR2004
                groups[-1].append(e)
            else:
                groups.append([e])
        if len(groups) < 2:  # noqa: PLR2004
            return [texts]
        spans = [(min(e["y_mm"] for e in g), max(e["y_mm"] for e in g)) for g in groups]
        overlap = any(
            spans[i][0] < spans[j][1] and spans[j][0] < spans[i][1]
            for i in range(len(spans))
            for j in range(i + 1, len(spans))
        )
        return groups if overlap else [texts]

    def _header_flow(
        self, container: Any, elements: list[Json], page_top: float, origin: float
    ) -> float:
        """Une pile de l'en-tête : paragraphes à leur hauteur mesurée ; rend le bas de la pile."""
        ordered = sorted(elements, key=lambda e: (round(e["y_mm"], 0), e["x_mm"]))
        rows: list[list[Json]] = []
        for e in ordered:
            st = self.styles[e["style"]]
            prev = rows[-1][-1] if rows else None
            same_line = (
                prev is not None
                and e["x_mm"] > prev["x_mm"] + 5
                and abs(prev["y_mm"] - e["y_mm"])
                < 0.5 * min(st["box_h_mm"], self.styles[prev["style"]]["box_h_mm"])
            )
            if same_line and rows:
                rows[-1].append(e)
            else:
                rows.append([e])
        prev_bottom = page_top
        first = True
        for row in rows:
            lead = row[0]
            st = self.styles[lead["style"]]
            # La première ligne de chaque pile part du haut de la page (marge) ;
            # une pile plus basse que la première garde son écart.
            before = lead["y_mm"] - prev_bottom if not first else lead["y_mm"] - page_top
            align = "center" if lead.get("align") == "center" else "left"
            indent = 0.0 if align == "center" else lead["x_mm"] - origin
            p = self.para(
                container,
                lead["style"],
                self.view["header"][lead["field"]],
                before_mm=before + (self.lead(lead["style"]) if first else 0.0),
                indent_mm=indent,
                align=align,
            )
            if self.anchor_paragraph is None and container is self.doc:
                self.anchor_paragraph = p
            first = False
            if lead.get("icon"):
                self.inline_icon(p, lead["icon"], before=True)
            for other in row[1:]:
                p.paragraph_format.tab_stops.add_tab_stop(Mm(other["x_mm"] - origin - indent))
                p.add_run("\t")
                run = p.add_run(self.view["header"][other["field"]])
                self.apply_run_style(run, other["style"])
            prev_bottom = lead["y_mm"] + float(st["box_h_mm"])
        return prev_bottom

    def apply_run_style(self, run: Any, key: str) -> None:
        """Le style d'un rôle appliqué à un morceau de paragraphe (deux champs sur une ligne)."""
        st = self.styles.get(key)
        if not st:
            return
        family, bold, italic = self.names.word_font(
            st["family"], int(st["weight"]), bool(st["italic"])
        )
        run.font.name = family
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), family)
        scale = 1.0 if key.startswith("header.") else self.f
        run.font.size = Pt(round(st["size_pt"] * scale * 2) / 2)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb(st["color"])
        run.font.all_caps = bool(st["uppercase"])
        if st["letter_spacing_pt"]:
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(round(st["letter_spacing_pt"] * 20)))
            run._r.get_or_add_rPr().append(sp)

    def inline_icon(self, p: Paragraph, icon: Json, *, before: bool) -> None:
        """Un pictogramme flottant à gauche du texte qu'il illustre, ancré au paragraphe."""
        data = self.spec.get("_assets", {}).get(icon["file"])
        if not data:
            return
        run = p.runs[0] if p.runs and before else p.add_run()
        holder = p.add_run()
        picture = holder.add_picture(
            io.BytesIO(data), width=Mm(icon["w_mm"]), height=Mm(icon["h_mm"])
        )
        inline = holder._r.xpath(".//wp:inline")[-1]
        extent, graphic, doc_pr = (
            inline.find(qn("wp:extent")),
            inline.find(qn("a:graphic")),
            inline.find(qn("wp:docPr")),
        )
        st_key = p.style.name if p.style is not None else ""
        del st_key, run, picture
        left = p.paragraph_format.left_indent or 0
        x = (left / EMU_PER_MM if isinstance(left, int) else 0.0) + icon["dx_mm"]
        anchor = parse_xml(
            '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="3" behindDoc="0" locked="0" '
            'layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
            f'<wp:positionH relativeFrom="column"><wp:posOffset>{emu(x)}</wp:posOffset></wp:positionH>'
            f'<wp:positionV relativeFrom="paragraph"><wp:posOffset>{emu(icon["dy_mm"] + self._para_lead(p))}</wp:posOffset></wp:positionV>'
            "</wp:anchor>"
        )
        anchor.append(extent)
        anchor.append(
            parse_xml(
                '<wp:effectExtent xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" l="0" t="0" r="0" b="0"/>'
            )
        )
        anchor.append(
            parse_xml(
                '<wp:wrapNone xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"/>'
            )
        )
        anchor.append(doc_pr)
        anchor.append(
            parse_xml(
                '<wp:cNvGraphicFramePr xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"/>'
            )
        )
        anchor.append(graphic)
        inline.getparent().replace(inline, anchor)

    def _para_lead(self, p: Paragraph) -> float:
        """Le blanc posé au-dessus du texte d'un paragraphe (espace avant + excédent d'interligne)."""
        name = p.style.name if p.style is not None else ""
        key = next((k for k, v in self._word_styles.items() if v == name), "section_title")
        before = p.paragraph_format.space_before
        before_mm = before / EMU_PER_MM if isinstance(before, int) else 0.0
        return before_mm + self.lead(key)

    # ------------------------------------------------------------------ bandes et rubriques
    def bands(self) -> None:
        """Chaque bande : un espaceur puis un tableau sans bordure (colonnes, gouttières)."""
        bands = placement(self.spec, self.view, self.fonts)
        for bi, band in enumerate(bands):
            if not any(band["columns"]):
                continue
            cols = band["spec"]["columns"]
            widths: list[tuple[str, float]] = []
            prev_right = 0.0
            for c in cols:
                widths.append(("gap", c["x_mm"] - prev_right))
                widths.append(("col", c["w_mm"]))
                prev_right = c["x_mm"] + c["w_mm"]
            if bi:
                gap_before = self.spec["bands"][bi - 1].get("gap_after_mm", 8.0) * self.k
            else:
                gap_before = self.spec["first_band_top_mm"] - self.header_bottom
            title_pitch = self.pitch(self.styles["section_title"])
            col_extra = [
                max(0.0, (col[0][0].get("title_pitch_mm") or title_pitch) - title_pitch)
                if col
                else 0.0
                for col in band["columns"]
            ]
            self._band_extra = (max(col_extra), col_extra)
            spacer = self.doc.add_paragraph()
            if self.first_spacer is None:
                self.first_spacer = spacer
            spacer_h = 0.35
            # Le titre de la première rubrique de chaque colonne part sans espace
            # avant : son texte commence un blanc d'interligne sous le haut de la cellule.
            spacer.paragraph_format.space_before = Mm(
                max(0.0, gap_before - spacer_h - self.lead("section_title") - max(col_extra))
            )
            spacer.paragraph_format.space_after = Pt(0)
            spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            spacer.paragraph_format.line_spacing = Mm(spacer_h)
            for d in self.spec["page_decorations"]:
                if d.get("anchor") == "band" and d["band"] == bi - 1:
                    spacer.add_run()._r.append(
                        anchored_shape(
                            x_mm=max(d["x_mm"], 0.0),
                            y_mm=d["dy_mm"] * self.k,
                            w_mm=d["w_mm"] + min(d["x_mm"], 0.0),
                            h_mm=d["h_mm"],
                            relative_h="page",
                            relative_v="paragraph",
                            kind=d["kind"],
                            fill=d.get("fill"),
                            stroke=d.get("stroke"),
                            stroke_pt=float(d.get("stroke_pt") or 0.75),
                        )
                    )
            table = self.doc.add_table(rows=1, cols=len(widths))
            no_borders(table)
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            set_widths(table, [max(w, 0.5) for _, w in widths])
            ci = 0
            for i, (kind, _) in enumerate(widths):
                if kind != "col":
                    continue
                cell = table.rows[0].cells[i]
                cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
                for n, (s, sec, _) in enumerate(band["columns"][ci]):
                    # Les premières lignes de toutes les colonnes s'alignent : une
                    # colonne dont le titre a un interligne plus grand part plus haut.
                    lift = max(col_extra) - col_extra[ci] if n == 0 else 0.0
                    self.section(cell, s, sec, cols[ci]["w_mm"], first=n == 0, lift=lift)
                if not cell.paragraphs:
                    cell.add_paragraph()
                ci += 1

    def section(  # noqa: PLR0912, PLR0915 - le gabarit d'une rubrique, dans l'ordre du modèle
        self,
        cell: _Cell,
        s: Json,
        sec: Json,
        col_w: float,
        *,
        first: bool,
        lift: float = 0.0,
    ) -> None:
        """Une rubrique : titre, éléments, puces, frise et pastilles."""
        kind = sec["kind"]
        margin = 0.0 if first else s["margin_top_mm"]
        title = self.para(
            cell, "section_title", sec["title"], before_mm=margin, indent_mm=s["title_dx_mm"]
        )
        if first:
            title.paragraph_format.space_before = Mm(lift)
        if s.get("title_w_mm"):
            title.paragraph_format.right_indent = Mm(
                max(0.0, col_w - s["title_dx_mm"] - s["title_w_mm"])
            )
        if s.get("title_pitch_mm"):
            # Interligne exact : l'excédent passe au-dessus de chaque ligne,
            # la première ligne garde sa place en rendant cet excédent sur l'espace avant.
            extra = s["title_pitch_mm"] - self.pitch(self.styles["section_title"])
            title.paragraph_format.line_spacing = Mm(s["title_pitch_mm"])
            before = title.paragraph_format.space_before
            before_mm = before / EMU_PER_MM if isinstance(before, int) else 0.0
            title.paragraph_format.space_before = Mm(max(0.0, before_mm - extra))
        for d in s["decorations"]:
            if d["anchor"] == "section":
                height = d["h_mm"]
                if d.get("stretch"):
                    height = max(
                        d["h_mm"],
                        estimate(self.styles, s, sec, col_w, self.fonts)
                        - s["margin_top_mm"]
                        - d["dy_mm"],
                    )
                self.ornament(title, d, d["dx_mm"], d["dy_mm"] + self.lead("section_title"), height)
        slots, order = extra_contact_slots(s, sec)
        for i, item in enumerate(sec["items"]):
            gap = s["first_gap_mm"] if i == 0 else s["item_gap_mm"]
            if i and kind in LIST_KINDS and gap < 1.0:
                gap = s["first_gap_mm"]
            first_row = True
            item_first: Paragraph | None = None
            item_paras: list[tuple[Paragraph, str]] = []
            for key in order:
                sl = slots[key]
                roles = sl["roles"]
                rows: list[tuple[list[tuple[str, str]], float]] = []
                if len(roles) == 1:
                    val = item.get(roles[0])
                    vals = val if isinstance(val, list) else ([val] if val else [])
                    for j, v in enumerate(vals):
                        repeat = (
                            sl["repeat_margin_mm"]
                            if sl["repeat_margin_mm"] is not None
                            else sl["margin_top_mm"]
                        )
                        rows.append(([(roles[0], v)], sl["margin_top_mm"] if j == 0 else repeat))
                else:
                    pairs = [(r, v) for r in roles if isinstance(v := item.get(r), str) and v]
                    if pairs:
                        rows.append((pairs, sl["margin_top_mm"]))
                for vals, m in rows:
                    before = gap if first_row else m
                    first_row = False
                    role0 = vals[0][0]
                    key0 = style_key(self.styles, kind, role0)
                    if len(vals) == 1:
                        p = self.list_para(cell, s, sl, key0, role0, text=vals[0][1], before=before)
                    else:
                        p = self.joined_para(cell, kind, sl, key0, vals=vals, before=before)
                    if kind == "contact" and role0 in s.get("icons", {}):
                        self.inline_icon(p, s["icons"][role0], before=True)
                    if item_first is None:
                        item_first = p
                    item_paras.append((p, role0))
            if item_first is not None:
                self.item_ornaments(
                    s, item_paras, last_item=i == len(sec["items"]) - 1, col_w=col_w
                )
        last = cell.paragraphs[-1] if cell.paragraphs else title
        for d in s["decorations"]:
            if d["anchor"] == "after":
                # Sous la rubrique (filet de séparation) : sous la dernière ligne
                # du dernier paragraphe, à l'écart mesuré.
                _, lines, pitch, _ = self.para_box(last, col_w)
                self.ornament(last, d, d["dx_mm"], lines * pitch + d["dy_mm"] * self.k, d["h_mm"])

    def list_para(
        self, cell: _Cell, s: Json, sl: Json, key0: str, role0: str, *, text: str, before: float
    ) -> Paragraph:
        """Un paragraphe d'une ligne de rubrique ; une puce du modèle devient un taquet + « • »."""
        mk = s.get("marker")
        if role0 in ROLE_MARKER and mk:
            p = self.para(
                cell, key0, text, before_mm=before, indent_mm=sl["dx_mm"], align=sl["align"]
            )
            self.bullet(p, key0, mk, sl["dx_mm"])
            return p
        return self.para(
            cell, key0, text, before_mm=before, indent_mm=sl["dx_mm"], align=sl["align"]
        )

    def bullet(self, p: Paragraph, key: str, mk: Json, text_x_mm: float) -> None:
        """Une vraie puce Word (liste à puces) : la personne ajoute une mission d'un retour à la ligne.

        La puce du modèle devient le glyphe « • » (« ■ » pour une puce carrée)
        de la police du texte, à la couleur du modèle, avec un retrait suspendu
        réglé pour que son encre commence là où commence la puce du modèle. Son
        corps reste celui du texte : avec un interligne exact, un glyphe plus
        grand (ou décalé verticalement) abaisse la ligne de base et LibreOffice
        rogne le haut de la ligne (constaté).
        """
        st = self.styles[key]
        glyph = "■" if mk["shape"] == "square" else "•"
        font = self.fonts.get(st["family"], int(st["weight"]), bool(st["italic"]))
        box = glyph_box(font.data, glyph) if font else None
        x0 = box[0] if box else 0.08
        size_pt = st["size_pt"] * self.f
        pen_x = text_x_mm + mk["dx_mm"] - x0 * size_pt / MM_TO_PT
        hanging = max(0.5, text_x_mm - pen_x)
        num_id = self.numbering(
            glyph, st, mk["color"], size_pt=size_pt, left_mm=text_x_mm, hanging_mm=hanging
        )
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num = OxmlElement("w:numId")
        num.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num)
        p_pr.insert(1 if p_pr.find(qn("w:pStyle")) is not None else 0, num_pr)
        p.paragraph_format.left_indent = Mm(text_x_mm)
        p.paragraph_format.first_line_indent = Mm(-hanging)

    def numbering(
        self, glyph: str, st: Json, color: str, *, size_pt: float, left_mm: float, hanging_mm: float
    ) -> int:
        """Une définition de puces par apparence (police, corps, couleur, retraits) ; réutilisée."""
        family, _, _ = self.names.word_font(st["family"], int(st["weight"]), False)
        key = (glyph, family, color, size_pt, round(left_mm, 2), round(hanging_mm, 2))
        if key in self._numbering:
            return self._numbering[key]
        numbering = self.doc.part.numbering_part.element
        abstract_id = 900 + len(self._numbering)
        num_id = 900 + len(self._numbering)
        abstract = parse_xml(
            f'<w:abstractNum xmlns:w="{W_NS}" w:abstractNumId="{abstract_id}">'
            '<w:multiLevelType w:val="singleLevel"/>'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
            f'<w:lvlText w:val="{glyph}"/><w:lvlJc w:val="left"/>'
            f'<w:pPr><w:tabs><w:tab w:val="num" w:pos="{twips(left_mm)}"/></w:tabs>'
            f'<w:ind w:left="{twips(left_mm)}" w:hanging="{twips(hanging_mm)}"/></w:pPr>'
            f"<w:rPr><w:rFonts w:ascii={quoteattr(family)} w:hAnsi={quoteattr(family)} w:cs={quoteattr(family)}/>"
            f'<w:color w:val="{color.lstrip("#")}"/><w:sz w:val="{round(size_pt * 2)}"/></w:rPr>'
            "</w:lvl></w:abstractNum>"
        )
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            numbering.append(abstract)
        numbering.append(
            parse_xml(
                f'<w:num xmlns:w="{W_NS}" w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/></w:num>'
            )
        )
        self._numbering[key] = num_id
        return num_id

    def joined_para(
        self,
        cell: _Cell,
        kind: str,
        sl: Json,
        key0: str,
        *,
        vals: list[tuple[str, str]],
        before: float,
    ) -> Paragraph:
        """Deux champs sur une ligne : taquet à l'écart mesuré, séparateur « | » à la couleur du filet."""
        p = self.para(cell, key0, vals[0][1], before_mm=before, indent_mm=sl["dx_mm"])
        sep = sl.get("separator")
        st0 = self.styles[key0]
        space_mm = self.fonts.width_mm(" ", {**st0, "letter_spacing_pt": 0.0}) * self.f
        for role, v in vals[1:]:
            if sep:
                # « | » à la couleur du filet, entouré des blancs mesurés : chaque
                # espace est élargi (interlettrage) jusqu'à l'écart du modèle.
                bar_mm = (
                    self.fonts.width_mm("|", {**st0, "letter_spacing_pt": 0.0, "uppercase": False})
                    * self.f
                )
                self._spaced_run(p, " ", sep["left_mm"] - space_mm - bar_mm / 2)
                glue = p.add_run("|")
                glue.font.color.rgb = rgb(sep["stroke"])
                glue.font.bold = False
                self._spaced_run(p, " ", sep["right_mm"] - space_mm - bar_mm / 2)
            else:
                self._spaced_run(p, " ", (sl["gap_mm"] or 3.0) - space_mm)
            run = p.add_run(v)
            self.apply_run_style(run, style_key(self.styles, kind, role))
        return p

    @staticmethod
    def _spaced_run(p: Paragraph, text: str, extra_mm: float) -> None:
        """Un blanc élargi de `extra_mm` (interlettrage du run, en vingtièmes de point)."""
        run = p.add_run(text)
        if abs(extra_mm) > 0.05:  # noqa: PLR2004
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), str(round(extra_mm * MM_TO_PT * 20)))
            run._r.get_or_add_rPr().append(sp)

    def wrap(self, text: str, key: str, width_mm: float) -> int:
        """Nombre de lignes d'un paragraphe : retour à la ligne glouton, mot à mot, police réelle."""
        st = self.styles.get(key) or self.styles["section_title"]
        scale = 1.0 if key.startswith("header.") else self.f
        style = {**st, "size_pt": st["size_pt"] * scale}
        space = self.fonts.width_mm(" ", {**style, "uppercase": False})
        lines, current = 1, 0.0
        for word in text.split():
            w = self.fonts.width_mm(word, style)
            if current and current + space + w > width_mm:
                lines += 1
                current = w
            else:
                current = current + space + w if current else w
        return lines

    def para_box(self, p: Paragraph, col_w: float) -> tuple[float, int, float, float]:
        """(espace avant, lignes, interligne, boîte d'une ligne) d'un paragraphe, en mm."""
        before, key, text, indent = self._geometry.get(id(p._p), (0.0, "section_title", "", 0.0))
        st = self.styles.get(key) or self.styles["section_title"]
        scale = 1.0 if key.startswith("header.") else self.f
        lines = self.wrap(text, key, max(10.0, col_w - indent))
        return before, lines, self.pitch(st) * scale, float(st["box_h_mm"]) * scale

    def item_ornaments(
        self,
        s: Json,
        paras: list[tuple[Paragraph, str]],
        *,
        last_item: bool,
        col_w: float,
    ) -> None:
        """Frise et pastilles d'un élément, ancrées à la première ligne de son premier paragraphe.

        La hauteur de l'élément se calcule comme Word posera le texte : chaque
        paragraphe occupe son espace avant puis ses lignes × l'interligne exact,
        l'excédent d'interligne au-dessus du texte. Une frise étirée s'arrête où
        l'élément s'arrête ; une pastille centrée sur un intitulé de deux lignes
        se centre sur les deux.
        """
        if not paras:
            return
        first_p, _ = paras[0]
        boxes = {id(p._p): self.para_box(p, col_w) for p, _ in paras}
        before0, _, pitch0, box0 = boxes[id(first_p._p)]
        item_top = before0 + (pitch0 - box0)
        item_bottom = sum(before + lines * pitch for before, lines, pitch, _ in boxes.values())
        item_h = item_bottom - item_top
        for d in s["decorations"]:
            if d["anchor"] != "item":
                continue
            if d.get("attach") == "row":
                target = next((p for p, r in paras if r == d.get("ref_role")), first_p)
                _, lines, pitch, box = boxes[id(target._p)]
                block = (lines - 1) * pitch + box
                top = (pitch - box) + block / 2 + d["center_offset_mm"] - d["h_mm"] / 2
                self.ornament(target, d, d["dx_mm"], top, d["h_mm"])
                continue
            inset = d.get("bottom_inset_mm", 0.0)
            if d.get("connector") and last_item:
                if d.get("last_bottom_inset_mm") is None:
                    continue
                inset = d["last_bottom_inset_mm"]
            height = d["h_mm"]
            if d.get("stretch"):
                height = max(1.0, item_h - d["dy_mm"] - inset * self.k)
            self.ornament(first_p, d, d["dx_mm"], d["dy_mm"] + (pitch0 - box0), height)

    @staticmethod
    def _space_before_mm(p: Paragraph) -> float:
        before = p.paragraph_format.space_before
        return before / EMU_PER_MM if isinstance(before, int) else 0.0

    def ornament(self, p: Paragraph, d: Json, x_mm: float, y_mm: float, h_mm: float) -> None:
        """Une forme ancrée à la première ligne du paragraphe `p`, en mm depuis le haut
        de cette ligne et le bord de la colonne.

        Repère « ligne » plutôt que « paragraphe » : LibreOffice place le haut
        d'un paragraphe de plusieurs lignes à interligne exact autrement que
        Word (constaté : 2 mm d'écart) ; la ligne est la même pour les deux. Le
        run qui porte l'ancre est en tête du paragraphe, donc sur sa première ligne.
        """
        run = OxmlElement("w:r")
        run.append(
            anchored_shape(
                x_mm=x_mm,
                y_mm=y_mm,
                w_mm=d["w_mm"],
                h_mm=h_mm,
                relative_h="column",
                relative_v="line",
                kind=d["kind"],
                fill=d.get("fill"),
                stroke=d.get("stroke"),
                stroke_pt=float(d.get("stroke_pt") or 0.75),
            )
        )
        first_run = p._p.find(qn("w:r"))
        if first_run is not None:
            first_run.addprevious(run)
        else:
            p._p.append(run)

    # ------------------------------------------------------------------ document
    def render(self, assets: dict[str, bytes]) -> bytes:
        """Le DOCX complet, polices incorporées."""
        self.spec = {**self.spec, "_assets": assets}
        sec = self.doc.sections[0]
        page = self.spec["page"]
        sec.page_width, sec.page_height = Mm(page["width_mm"]), Mm(page["height_mm"])
        texts = [e for e in self.spec["header"]["elements"] if self.view["header"].get(e["field"])]
        top = min([e["y_mm"] for e in texts], default=10.0)
        st_top = self.styles[min(texts, key=lambda e: e["y_mm"])["style"]] if texts else None
        lead_top = (self.pitch(st_top) - st_top["box_h_mm"]) if st_top else 0.0
        sec.top_margin, sec.bottom_margin = Mm(max(0.0, top - lead_top)), Mm(6)
        sec.left_margin, sec.right_margin = Mm(0), Mm(0)
        sec.header_distance, sec.footer_distance = Mm(0), Mm(0)
        normal = self.doc.styles["Normal"]
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.space_before = Pt(0)
        normal.font.size = Pt(2)
        self.header()
        self.bands()
        # Les formes de page s'ancrent hors des tableaux (dans une cellule,
        # Word les placerait par rapport à la cellule).
        anchor = self.anchor_paragraph or self.first_spacer
        if anchor is not None:
            self.page_shapes(anchor)
        props = self.doc.core_properties
        full_name = self.view["header"].get("full_name", "")
        props.title = f"CV — {full_name}".strip(" —")
        props.author = full_name
        props.keywords = self.view["header"].get("headline", "")
        buffer = io.BytesIO()
        self.doc.save(buffer)
        return embed_fonts(buffer.getvalue(), self.names)
